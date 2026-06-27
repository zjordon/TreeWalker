"""Action execution engine — registers all browser actions and executes them."""

from __future__ import annotations

import asyncio
import json
from typing import Any
import logging
import mimetypes
import os
import re
import shutil
import time

from pydantic import BaseModel, ValidationError

from tree_walker.agent.views import ActionResult
from tree_walker.browser.session import BrowserSession, _requires_direct_value_assignment
from tree_walker.browser.views import BrowserStateSummary, EnhancedDOMTreeNode, SerializedDOMState
from tree_walker.config import TruncationSettings
from tree_walker.tools.extract_markdown import chunk_markdown_by_structure, extract_clean_markdown
from tree_walker.tools.models import ACTION_DEFINITIONS, make_structured_done_params
from tree_walker.tools.registry import ActionRegistry

logger = logging.getLogger(__name__)

# G5 进阶：空选项按下拉类型给诊断提示（native 不在此表 -> 沿用 P0 仅 hint，零回归）。
# click-select 也列入：click SELECT 分支空选项时给同样的懒加载提示。
_EMPTY_OPTIONS_DIAGNOSTIC = {
	"aria": "Listbox/menu found but no [role=option] children (may need expanding).",
	"custom": "Custom dropdown found but no .item/.option/[data-value] children (may need expanding).",
	"combobox": "Combobox listbox found but empty (options may load on demand).",
	"click-select": "Select element has no <option> children (may populate lazily — try dropdown_options again after the page settles).",
}


# ── Helpers ───────────────────────────────────────────────────────────


def _is_file_input_node(node: EnhancedDOMTreeNode) -> bool:
    return node.tag_name.upper() == "INPUT" and node.attributes.get("type", "").lower() == "file"


def _file_matches_accept(file_path: str, accept: str | None) -> bool:
    """True if file_path is acceptable under the input's ``accept`` attribute.

    Parses the HTML accept attribute (comma-separated tokens, each a file
    extension like ``.png``, a wildcard MIME like ``image/*``, or a full MIME
    like ``application/pdf``) and matches the file's extension / guessed MIME.
    Empty/missing accept means "no restriction" (True).

    Uses stdlib mimetypes to map the extension to a MIME, so wildcard and
    full-MIME tokens work without a hard-coded table. Used by
    Tools._action_upload_file to emit a soft ⚠️ Note on mismatch — it never
    blocks the upload (CDP / browser-use do not either).
    """
    accept = (accept or "").strip()
    if not accept:
        return True
    file_ext = os.path.splitext(file_path)[1].lower()
    guessed_mime, _ = mimetypes.guess_type(file_path)
    for token in accept.split(","):
        token = token.strip().lower()
        if not token:
            continue
        if token.startswith("."):
            if file_ext == token:
                return True
        elif token.endswith("/*"):
            prefix = token[:-1]  # "image/"
            if guessed_mime and guessed_mime.startswith(prefix):
                return True
        else:
            if guessed_mime == token:
                return True
    return False


def _sniff_file_kind(path: str) -> str:
    """Peek magic bytes (阶段二 二.B) to classify a file before decoding.

    Returns one of ``"text"``, ``"pdf"``, ``"docx"``, ``"image"``, ``"binary"``.
    Used by ``_action_read_file`` to reject unsupported binaries early (actionable
    error instead of a cryptic UnicodeDecodeError) and to dispatch rich docs
    (PDF/DOCX/image → ``_read_rich_document``). Magic bytes are primary;
    extension is secondary because zip-family containers (docx/xlsx/pptx/plain
    zip) share ``PK\\x03\\x04`` and WebP shares ``RIFF`` with AVI/WAV.
    """
    with open(path, "rb") as f:
        head = f.read(12)
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image"
    if head.startswith(b"\xff\xd8\xff"):  # JPEG
        return "image"
    if head.startswith(b"GIF87a") or head.startswith(b"GIF89a"):
        return "image"
    if head.startswith(b"RIFF"):
        return "image" if head[8:12] == b"WEBP" else "binary"  # WebP vs AVI/WAV
    if head.startswith(b"%PDF-"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):  # zip container
        return "docx" if path.lower().endswith(".docx") else "binary"
    if head.startswith(b"\x7fELF") or head.startswith(b"MZ"):  # executables
        return "binary"
    if head.startswith(b"\x1f\x8b") or head.startswith(b"BZh") or head.startswith(b"Rar!"):
        return "binary"  # gzip / bzip2 / rar
    if head.startswith(b"7z\xbc\xaf\x27\x1c"):
        return "binary"  # 7z
    return "text"


def _format_search_results(data: dict, query: str) -> str:
    """Format search_page {matches, total, has_more, offset, attribute_matches}
    as LLM-readable text.

    Mirrors browser-use ``_format_search_results`` (service.py:337-360), with
    Phase 2 additions: an offset-aware pagination footer and a separate
    attribute-matches section. Caller guarantees total > 0 or attribute_total >
    0 (the total==0 soft-miss path is handled in _action_search_page).
    """
    matches = data.get("matches", [])
    total = data.get("total", 0)
    has_more = data.get("has_more", False)
    offset = data.get("offset", 0)

    lines = [f'Found {total} match{"es" if total != 1 else ""} for "{query}" on page:', ""]
    for i, m in enumerate(matches):
        context = m.get("context", "")
        path = m.get("element_path", "")
        loc = f" (in {path})" if path else ""
        lines.append(f"[{i + 1}] {context}{loc}")
    if has_more:
        next_offset = offset + len(matches)
        lines.append(
            f"\n... showing {offset + 1}–{offset + len(matches)} of {total} total matches. "
            f"Call again with offset={next_offset} for the next batch (or raise max_results)."
        )
    attr_matches = data.get("attribute_matches") or []
    attr_total = data.get("attribute_total", 0)
    if attr_total:
        lines.append("")
        lines.append(f'Attribute matches for "{query}" ({attr_total}):')
        for i, m in enumerate(attr_matches):
            path = m.get("element_path", "")
            loc = f" (in {path})" if path else ""
            lines.append(f"[{i + 1}] @{m.get('attribute', '')}={m.get('value', '')}{loc}")
        if attr_total > len(attr_matches):
            lines.append(f"... showing {len(attr_matches)} of {attr_total} attribute matches.")
    return "\n".join(lines)


def _format_find_results(data: dict, selector: str) -> str:
    """Format find_elements {elements, total, showing, offset, has_more} as
    LLM-readable text.

    Mirrors browser-use ``_format_find_results`` (``service.py:363-401``), with
    Phase 2 additions: per-element origin tag (shadow DOM / iframe), optional
    geometry+visibility, and an offset-aware pagination footer (mirrors
    ``_format_search_results``). Caller guarantees total > 0 (the total==0
    soft-miss path is handled in _action_find_elements).
    """
    elements = data.get("elements", [])
    total = data.get("total", 0)
    offset = data.get("offset", 0)
    has_more = data.get("has_more", False)

    lines = [f'Found {total} element{"s" if total != 1 else ""} matching "{selector}":', ""]
    for el in elements:
        idx = el.get("index", 0)
        tag = el.get("tag", "?")
        text = el.get("text", "")
        attrs = el.get("attrs", {})
        children = el.get("children_count", 0)
        origin = el.get("origin", "")
        rect = el.get("rect")

        parts = [f"[{idx}] <{tag}>"]
        if text:
            display_text = " ".join(text.split())
            if len(display_text) > 120:
                display_text = display_text[:120] + "..."
            parts.append(f'"{display_text}"')
        if attrs:
            attr_strs = [f'{k}="{v}"' for k, v in attrs.items()]
            parts.append("{" + ", ".join(attr_strs) + "}")
        parts.append(f"({children} children)")
        if rect:
            vis = "visible" if el.get("visible") else "hidden"
            parts.append(f"({vis}, {int(rect['w'])}x{int(rect['h'])}@{int(rect['x'])},{int(rect['y'])})")
        if origin:
            parts.append(origin.strip())          # ' (in shadow DOM)' → 'in shadow DOM'
        lines.append(" ".join(parts))

    if has_more:
        next_offset = offset + len(elements)
        lines.append(
            f"\n... showing {offset + 1}–{offset + len(elements)} of {total} total elements. "
            f"Call again with offset={next_offset} for the next batch (or raise max_results)."
        )
    return "\n".join(lines)


def _format_node_id_results(data: dict, selector: str) -> str:
    """Format find_elements_node_ids {node_ids, total, showing, offset, has_more}
    as LLM-readable text.

    Each ``backend_id`` is a backendNodeId usable directly as the
    ``index``/``element_id`` of click/input_text (index===backend_id in this
    system). Caller guarantees total > 0 (the total==0 soft-miss path is
    handled in _action_find_elements).
    """
    node_ids = data.get("node_ids", [])
    total = data.get("total", 0)
    offset = data.get("offset", 0)
    has_more = data.get("has_more", False)

    lines = [f'Found {total} element{"s" if total != 1 else ""} matching "{selector}" (node ids):', ""]
    for el in node_ids:
        bid = el.get("backend_id")
        tag = el.get("tag", "?")
        lines.append(f"[{bid}] <{tag}>  (pass as index= or element_id= to click/input_text)")
    if has_more:
        next_offset = offset + len(node_ids)
        lines.append(
            f"\n... showing {offset + 1}–{offset + len(node_ids)} of {total} total elements. "
            f"Call again with offset={next_offset} for the next batch."
        )
    return "\n".join(lines)


_EVAL_MEMORY_ECHO_MAX = 200


def _eval_long_term_memory(text: str) -> str:
    """Build a compact long_term_memory for an evaluate result.

    Short results are echoed verbatim (so ``return document.title`` is
    remembered across steps); long results collapse to a length-only summary
    to avoid bloating cross-step memory. Mirrors browser-use
    service.py:1847-1853 (their threshold is 10000; ours is tighter).
    """
    if len(text) <= _EVAL_MEMORY_ECHO_MAX:
        return text
    return f"JavaScript executed successfully, result length: {len(text)} characters."


_DATA_IMAGE_RE = re.compile(r"data:image/[a-zA-Z0-9.+-]+;base64,[A-Za-z0-9+/=]+")


def _extract_data_images(text: str) -> tuple[str, list[str]]:
    """Split ``data:image/...;base64,...`` URIs out of ``text`` (阶段二 二.F).

    Returns ``(text_with_placeholders, images)``: each URI is replaced by
    ``[image N]`` and ``images`` lists the extracted data URIs in order, so the
    bulky base64 stays out of ``extracted_content`` (moved to ``metadata``).
    Mirrors browser-use service.py:1821-1836, adapted to a free-text scan.
    """
    images: list[str] = []

    def _repl(m: re.Match) -> str:
        images.append(m.group(0))
        return f"[image {len(images)}]"

    return _DATA_IMAGE_RE.sub(_repl, text), images


def _find_upload_label_near(
    node: EnhancedDOMTreeNode, max_ancestor: int = 4, max_depth: int = 3,
) -> int | None:
    """Locate the nearest ``<label>`` upload trigger relative to ``node``.

    A ``<label for=input>`` click is dispatched by the BROWSER straight to its
    associated input — this native label semantics preserves the transient
    user-activation that opening a file chooser requires. By contrast, the JS
    ``input.click()`` that upload frameworks (Semi-UI / AntD Upload) fire from a
    drag-area or button click loses that activation (usually across an async
    gap), so no ``Page.fileChooserOpened`` ever fires. This is the crux of issue
    #34: on Douyin's cover editor, clicking the drag-area / a button never opens
    a chooser, but clicking the sibling ``<label class*="upload">`` (选择文件)
    does, and ``setFileInputFiles`` on the input it opens sets the cover cleanly
    (empirically verified).

    Such uploaders render the label as a sibling of the drag-area inside one
    Upload container, so we climb from the target to that container and return
    its upload label. Returns the label's ``backend_node_id`` or ``None``.
    """
    def _is_upload_label(n: EnhancedDOMTreeNode | None) -> bool:
        if n is None or (n.tag_name or "").upper() != "LABEL":
            return False
        cls = ((n.attributes or {}).get("class") or "").lower()
        parts = [(n.node_value or "")]
        for c in (n.children_nodes or [])[:6]:
            parts.append((c.node_value or "") if c is not None else "")
        txt = " ".join(p for p in parts if p)
        return "upload" in cls or "选择文件" in txt or "上传" in txt

    def _search_subtree(root, depth):
        if root is None or depth < 0:
            return None
        if _is_upload_label(root) and root.backend_node_id is not None:
            return root.backend_node_id
        for c in (root.children_nodes or []):
            found = _search_subtree(c, depth - 1)
            if found:
                return found
        for sr in (root.shadow_roots or []):
            found = _search_subtree(sr, depth - 1)
            if found:
                return found
        return None

    cur = node
    depth = 0
    while cur is not None and depth <= max_ancestor:
        found = _search_subtree(cur, max_depth)
        if found:
            return found
        cur = cur.parent_node
        depth += 1
    return None


# ── Search engines ──────────────────────────────────────────────────

_SEARCH_ENGINE_URLS: dict[str, str] = {
    "baidu": "https://www.baidu.com/s?wd={query}",
    "google": "https://www.google.com/search?q={query}&udm=14",
    "bing": "https://www.bing.com/search?q={query}",
    "duckduckgo": "https://duckduckgo.com/?q={query}",
}


# ── Navigate health check / error mapping (mirrors browser-use) ──────

# 等待/重试时长（秒），参照 browser_use/tools/service.py:501-523
_NAVIGATE_EMPTY_RETRY_WAIT = 3.0   # 首次发现空 DOM 后等待重查
_NAVIGATE_EMPTY_RELOAD_WAIT = 5.0  # reload 后等待

# 网络错误码 → 触发 "site unavailable" 友好提示（参照 browser-use service.py:544-557）
_NAVIGATE_NET_ERROR_MARKERS = (
    "ERR_NAME_NOT_RESOLVED",
    "ERR_INTERNET_DISCONNECTED",
    "ERR_CONNECTION_REFUSED",
    "ERR_TIMED_OUT",
    "ERR_TUNNEL_CONNECTION_FAILED",
    "net::",
)


class Tools:
    """Action registry + execution engine.

    Each browser action is registered as a closure bound to this instance.
    The registry produces the Anthropic tool schema used by the LLM client.
    """

    def __init__(
        self,
        truncation: TruncationSettings | None = None,
        allowed_upload_paths: list[str] | None = None,
        allowed_write_paths: list[str] | None = None,
        allowed_read_paths: list[str] | None = None,
        display_files_in_done_text: bool = False,
        output_model: type[BaseModel] | None = None,
    ) -> None:
        # 二.E：output_model 须在 _register_all 之前落到 self / registry（_register_all 据此选变体）。
        self._output_model = output_model
        self.registry = ActionRegistry(output_model=output_model)
        self._register_all()
        self._cached_browser_state: BrowserStateSummary | None = None
        self._truncation = truncation or TruncationSettings()
        self._allowed_upload_paths = allowed_upload_paths
        self._allowed_write_paths = allowed_write_paths
        self._allowed_read_paths = allowed_read_paths
        self._display_files_in_done_text = display_files_in_done_text

    # ── Public API ─────────────────────────────────────────────────────

    async def execute(
        self,
        action_name: str,
        params: dict,
        browser: BrowserSession,
        browser_state: BrowserStateSummary | None = None,
    ) -> ActionResult:
        """Execute a single action by name with given parameters."""
        if action_name not in self.registry.actions:
            return ActionResult(error=f"Unknown action: {action_name}")

        params = self._flatten_params(params, action_name)

        self._cached_browser_state = browser_state
        try:
            registered = self.registry.actions[action_name]
            result = await registered.handler(params, browser)
            return self._normalize(result)
        except Exception as e:
            logger.exception("Action %s failed", action_name)
            return ActionResult(error=str(e))
        finally:
            self._cached_browser_state = None

    # ── Element lookup ──────────────────────────────────────────────────

    async def _get_element_by_index(
        self, index: int, browser: BrowserSession,
    ) -> tuple[Any | None, ActionResult | None]:
        """Look up element by index, using cached state when available."""
        if self._cached_browser_state and self._cached_browser_state.dom_state:
            entry = self._cached_browser_state.dom_state.selector_map.get(index)
            if entry:
                return entry, None

        state = await browser.get_state(include_screenshot=False)
        if not state.dom_state:
            return None, ActionResult(error="No DOM state available")
        entry = state.dom_state.selector_map.get(index)
        if not entry:
            return None, ActionResult(error=f"Element {index} not found in DOM state")
        return entry, None

    # ── Registration ───────────────────────────────────────────────────

    def _register_all(self) -> None:
        for name, (param_model, description, terminates) in ACTION_DEFINITIONS.items():
            # 二.E：done 结构化输出——output_model 给定时用变体 B 参数模型（data: T）。
            if name == "done" and self._output_model is not None:
                param_model = make_structured_done_params(self._output_model)
            handler = getattr(self, f"_action_{name}", None)
            if handler is None:
                logger.debug("No handler for action %s, skipping", name)
                continue
            self.registry.action(
                name=name,
                description=description,
                param_model=param_model,
                terminates=terminates,
            )(handler)

    def apply_page_filters(self, filters: dict[str, list[str]]) -> None:
        """Apply page-pattern filters to specific actions.

        Each key is an action name; the value is a list of glob patterns.
        Unknown action names are silently ignored.
        """
        for name, patterns in filters.items():
            if name in self.registry.actions:
                self.registry.actions[name].page_patterns = patterns

    # ── Action handlers ────────────────────────────────────────────────

    async def _action_navigate(self, params: dict, browser: BrowserSession) -> ActionResult:
        url = params["url"]
        if not url.startswith(("http://", "https://")):
            url = "https://" + url
        new_tab = params.get("new_tab", False)

        try:
            await browser.navigate(url, new_tab=new_tab)
            # 健康检查：仅当前标签页 + http(s) URL（chrome://、about:、new_tab=True 跳过）
            if not new_tab:
                await self._navigate_health_check(url, browser)
            memory = (
                f"Opened new tab with URL {url}" if new_tab else f"Navigated to {url}"
            )
            logger.info(memory)
            return ActionResult(extracted_content=memory, long_term_memory=memory)
        except Exception as e:
            return self._map_navigation_error(url, e)

    @staticmethod
    def _dom_appears_empty(state: BrowserStateSummary) -> bool:
        """判定页面是否为空（镜像 browser-use 的 _page_appears_empty）。

        SerializedDOMState.llm_representation() 在 _root is None 时返回非空占位符，所以必须
        单独检查 _root is None（不能只看 element_tree_text）。
        """
        ds = state.dom_state
        if ds is None:
            return True
        return ds._root is None or not ds.element_tree_text.strip()

    async def _navigate_health_check(self, url: str, browser: BrowserSession) -> None:
        """导航后检测空 DOM，参照 browser_use/tools/service.py:493-523 三阶段重试。

        仅 http(s) URL + 当前标签页触发；三阶段判定逐渐收严。
        """
        state = await browser.get_state(include_screenshot=False)
        url_is_http = state.url.lower().startswith(("http://", "https://"))
        if not (url_is_http and self._dom_appears_empty(state)):
            return

        logger.warning(
            "Empty DOM after navigating to %s, waiting %.0fs and rechecking",
            url, _NAVIGATE_EMPTY_RETRY_WAIT,
        )
        await asyncio.sleep(_NAVIGATE_EMPTY_RETRY_WAIT)
        state = await browser.get_state(include_screenshot=False)
        if not (state.url.lower().startswith(("http://", "https://")) and self._dom_appears_empty(state)):
            return

        logger.warning("Still empty after %.0fs, reloading %s", _NAVIGATE_EMPTY_RETRY_WAIT, url)
        # reload：重新 navigate，异常吞掉（避免健康检查的二次失败中断"初次导航已成功"的外层路径）
        try:
            await browser.navigate(url, new_tab=False)
        except Exception as reload_err:
            logger.warning("Reload during health check failed: %s", reload_err)
        await asyncio.sleep(_NAVIGATE_EMPTY_RELOAD_WAIT)

        state = await browser.get_state(include_screenshot=False)
        if state.url.lower().startswith(("http://", "https://")):
            ds = state.dom_state
            if ds is None or ds._root is None:
                raise RuntimeError(
                    f"Page loaded but returned empty content for {url}. "
                    f"The page may require JavaScript that failed to render, use anti-bot measures, "
                    f"or have a connection issue (e.g. tunnel/proxy error). Try a different URL or approach."
                )

    @staticmethod
    def _map_navigation_error(url: str, e: Exception) -> ActionResult:
        """把导航异常映射为对 LLM 友好的 ActionResult.error（参照 browser-use service.py:534-560）。"""
        error_msg = str(e)
        if any(marker in error_msg for marker in _NAVIGATE_NET_ERROR_MARKERS):
            logger.warning("Navigation to %s failed - site unavailable: %s", url, error_msg)
            return ActionResult(error=f"Navigation failed - site unavailable: {url}")
        return ActionResult(error=f"Navigation failed: {error_msg}")

    async def _action_click(self, params: dict, browser: BrowserSession) -> ActionResult:
        # 0. element_id is an alias for index (a backend node id from
        # find_elements(return_node_ids=True); index===backend_id in this
        # system, so both resolve through selector_map). Exactly one required —
        # the registry does not validate the execute path, so guard here.
        index = params.get("index")
        element_id = params.get("element_id")
        if (index is None) == (element_id is None):
            return ActionResult(
                error="click requires exactly one of `index` or `element_id`.",
            )
        params = {**params, "index": index if index is not None else element_id}
        # 1. 元素查找（保持原逻辑）
        entry, error = await self._get_element_by_index(params["index"], browser)
        if error:
            return error

        backend_id = entry.backend_node_id

        # 1.5 file-input 守卫：点 <input type=file> 只会弹原生文件框（拦截已兜底不弹框，
        # 但 click 仍会派发到页面、且对上传无任何有用效果）。引导 LLM 改用 upload_file。
        if _is_file_input_node(entry):
            return ActionResult(
                error=(
                    f"Element {params['index']} is an <input type='file'>. Clicking it "
                    f"would open the OS file picker. Use upload_file(index={params['index']}, "
                    f"path=<absolute path>) instead to set the file programmatically."
                ),
            )

        # 2. 下拉降级（G9）：目标是下拉形元素（native select / ARIA listbox·menu / custom
        #    dropdown / combobox）时，click 要么 no-op（native select 忽略 click）要么误开
        #    遮罩（ARIA/combobox）。降级为读选项给 LLM（镜像 dropdown_options 调度，复用
        #    _format_options_result，输出与 dropdown_options 字节级同格式，修掉原 str(options)
        #    repr 不一致）。D7：不重路由到 dropdown_options（click 已「点过」）。非下拉落回真实 click。
        tag = entry.tag_name.upper()
        is_combo, _ = self._is_autocomplete_field(entry)
        attrs = getattr(entry, "attributes", {}) or {}
        is_dropdown_target = (
            tag == "SELECT"
            or attrs.get("role") in ("listbox", "menu", "menubar", "tree", "grid")
            or "dropdown" in (attrs.get("class") or "").lower()
            or (is_combo and (attrs.get("aria-controls") or attrs.get("aria-owns")))
        )
        if is_dropdown_target:
            try:
                if tag == "SELECT":
                    options = await browser.fetch_select_options(backend_id)
                    return self._format_options_result(options, entry, params["index"], "click-select")
                if is_combo and (attrs.get("aria-controls") or attrs.get("aria-owns")):
                    options = await browser.expand_and_fetch_combobox_options(backend_id)
                    return self._format_options_result(options, entry, params["index"], "click-combobox")
                dispatched = await browser.fetch_dropdown_options(backend_id)
                if dispatched["source"] is not None:
                    return self._format_options_result(
                        dispatched["options"], entry, params["index"], "click-" + dispatched["source"],
                    )
                # 假阳性：判错为下拉但 dispatcher 真阴性 → 落回真实 click
            except Exception as e:
                return ActionResult(error=f"Failed to read dropdown options: {e}")

        # 3. 普通点击：highlight -> click_element，映射 bool 信号
        tabs_before = tuple(t.target_id for t in await browser.get_tabs())  # G7 新页检测快照
        try:
            await browser.highlight_element(backend_id)
            clicked = await browser.click_element(backend_id)
        except Exception as e:
            # CDP 异常（连接断开、target 消失等）——友好映射，不让 LLM 看裸堆栈
            return ActionResult(error=f"Click failed: {e}")

        if not clicked:
            # 坐标拿不到 + JS 回退也失败 —— 明确告知 LLM，不再静默成功
            return ActionResult(
                error=(
                    f"Could not click element {params['index']} "
                    f"(no coordinates and JS click fallback failed; "
                    f"the element may be detached, hidden, or in a cross-origin iframe)"
                ),
            )

        # 4. 成功回显（对齐 navigate/go_back 风格）+ 新标签页检测（G7）
        memory = self._describe_click(entry, params["index"])
        memory += await self._detect_new_tab_opened(browser, tabs_before)
        logger.info(memory)
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    @staticmethod
    def _describe_click(entry: Any, index: int) -> str:
        """Build a human-readable click echo, mirroring navigate/go_back style.

        Prefers an identifying attribute the LLM can also see in the DOM tree
        (aria-label/placeholder/title/alt/value), then node_value, then just the
        tag. Bounded to ~60 chars per field so the echo fits the LLM context.
        """
        tag = entry.tag_name.upper()
        attrs = getattr(entry, "attributes", {}) or {}
        for key in ("aria-label", "placeholder", "title", "alt", "value"):
            v = attrs.get(key)
            if v:
                v = v.strip()
                if len(v) > 60:
                    v = v[:60] + "..."
                return f"Clicked [{tag}] {v!r} at index {index}"
        node_value = (getattr(entry, "node_value", "") or "").strip()
        if node_value:
            if len(node_value) > 60:
                node_value = node_value[:60] + "..."
            return f"Clicked [{tag}] {node_value!r} at index {index}"
        return f"Clicked [{tag}] at index {index}"

    async def _detect_new_tab_opened(
        self, browser: BrowserSession, tabs_before: tuple[str, ...],
    ) -> str:
        """点击若打开了新标签页，自动切过去并返回给 LLM 的提示串。

        对齐 browser-use _detect_new_tab_opened：click 前快照 target_id 集合，
        click 后 diff，命中新页则 switch_tab 并回显；切换失败软降级为提示。
        常态点击（未开新页）返回空串，不影响原有回显。
        """
        try:
            await asyncio.sleep(0.05)  # 等 Target.attachedToTarget 事件传播
            tabs_after = await browser.get_tabs()
            new_tabs = [t for t in tabs_after if t.target_id not in tabs_before]
            if not new_tabs:
                return ""
            new_tab = new_tabs[0]
            new_id = new_tab.target_id[-4:]
            try:
                await browser.switch_tab(new_tab.target_id)
                return (f"  ℹ️ Click opened a new tab [{new_id}] {new_tab.title}; "
                        f"auto-switched to it.")
            except Exception:
                return (f"  ℹ️ Click opened a new tab [{new_id}] {new_tab.title}; "
                        f"use switch_tab to focus it.")
        except Exception:
            return ""

    @staticmethod
    def _describe_input(entry: Any, index: int, text: str) -> str:
        """Build a human-readable input echo, mirroring _describe_click /
        navigate / go_back style.

        Prefers an identifying attribute the LLM can also see in the DOM tree
        (aria-label/placeholder/title), then node_value, then just the tag.
        Both the label and the typed text are bounded to ~60 chars so the echo
        fits the LLM context. Skips 'value'/'alt' (unlike _describe_click) since
        for an input the typed text itself is what matters.
        """
        shown = text if len(text) <= 60 else text[:60] + "..."
        tag = entry.tag_name.upper()
        attrs = getattr(entry, "attributes", {}) or {}
        for key in ("aria-label", "placeholder", "title"):
            v = attrs.get(key)
            if v:
                v = v.strip()
                if len(v) > 60:
                    v = v[:60] + "..."
                return f"Typed {shown!r} into [{tag}] {v!r} at index {index}"
        node_value = (getattr(entry, "node_value", "") or "").strip()
        if node_value:
            if len(node_value) > 60:
                node_value = node_value[:60] + "..."
            return f"Typed {shown!r} into [{tag}] {node_value!r} at index {index}"
        return f"Typed {shown!r} into [{tag}] at index {index}"

    @staticmethod
    def _describe_upload(entry: Any, index: int, file_path: str) -> str:
        """Build a human-readable upload echo, mirroring _describe_click /
        _describe_input / navigate / go_back style.

        Shows the uploaded file's basename (not the full path — which is long
        and noisy) plus an identifying attribute the LLM can also see in the
        DOM tree (aria-label/title/name/placeholder), then node_value, then
        just the tag. Skips 'value'/'alt': a file input's value is the
        browser-faked 'C:\\fakepath\\<name>' (not useful) and alt is irrelevant
        here. Bounded to ~60 chars per field so the echo fits the LLM context.
        """
        shown = os.path.basename(file_path)
        if len(shown) > 60:
            shown = shown[:60] + "..."
        tag = entry.tag_name.upper()
        attrs = getattr(entry, "attributes", {}) or {}
        for key in ("aria-label", "title", "name", "placeholder"):
            v = attrs.get(key)
            if v:
                v = v.strip()
                if len(v) > 60:
                    v = v[:60] + "..."
                return f"Uploaded {shown!r} to [{tag}] {v!r} at index {index}"
        node_value = (getattr(entry, "node_value", "") or "").strip()
        if node_value:
            if len(node_value) > 60:
                node_value = node_value[:60] + "..."
            return f"Uploaded {shown!r} to [{tag}] {node_value!r} at index {index}"
        return f"Uploaded {shown!r} to [{tag}] at index {index}"

    @staticmethod
    def _describe_dropdown(entry: Any, index: int) -> str:
        """Build a human-readable dropdown echo, mirroring _describe_click /
        _describe_upload style.

        select 元素通常没有 placeholder/value（value 是当前选中值），故优先级链用
        aria-label/title/name/id，再退到 node_value，最后退到 tag。Bounded to ~60 chars.
        """
        tag = (getattr(entry, "tag_name", "") or "").upper() or "SELECT"
        attrs = getattr(entry, "attributes", {}) or {}
        for key in ("aria-label", "title", "name", "id"):
            v = attrs.get(key)
            if v:
                v = v.strip()
                if len(v) > 60:
                    v = v[:60] + "..."
                return f"[{tag}] {v!r} at index {index}"
        node_value = (getattr(entry, "node_value", "") or "").strip()
        if node_value:
            if len(node_value) > 60:
                node_value = node_value[:60] + "..."
            return f"[{tag}] {node_value!r} at index {index}"
        return f"[{tag}] at index {index}"

    def _format_options_result(
        self, raw_options: list[dict], entry: Any, index: int, source: str,
    ) -> ActionResult:
        """Shared dropdown echo (short/long split) for native/aria/custom/combobox/
        subtree/click-select paths. ``source`` folds into long_term_memory as the
        subtree-search diagnostic channel; "native" renders no suffix
        (byte-identical to P0 so existing tests don't regress)."""
        lines: list[str] = []
        for i, opt in enumerate(raw_options):
            text = json.dumps(opt.get("text", ""))
            value = json.dumps(opt.get("value", ""))
            status = " (selected)" if opt.get("selected") else ""
            lines.append(f"{i}: text={text}, value={value}{status}")
        hint = f"Use the value in select_dropdown(index={index}, value=...)"
        # G5 进阶：空选项按 source 给诊断（native 沿用 P0 仅 hint，避免回归）
        if not lines:
            diag = _EMPTY_OPTIONS_DIAGNOSTIC.get(source)
            extracted = (diag + "\n" + hint) if diag else hint
        else:
            extracted = "\n".join(lines) + "\n" + hint
        desc = self._describe_dropdown(entry, index)
        if source == "native":
            via = ""
        elif source.startswith("child-depth-"):
            via = f" via {source}"
        else:
            via = f" via [{source.upper()}]"
        memory = f"Got {len(raw_options)} options from {desc}{via}"
        return ActionResult(extracted_content=extracted, long_term_memory=memory)

    @staticmethod
    def _find_node_by_backend_id(
        backend_node_id: int | None,
        dom_state: SerializedDOMState | None,
    ) -> EnhancedDOMTreeNode | None:
        """Look up the DOM node whose backend_node_id matches in the cached
        selector_map. Returns None if not found (e.g. the file input is hidden
        and excluded from the interactive selector_map). Used by
        _action_upload_file to read the resolved file input's ``accept``
        attribute for soft validation.
        """
        if not dom_state or backend_node_id is None:
            return None
        for node in dom_state.selector_map.values():
            if getattr(node, "backend_node_id", None) == backend_node_id:
                return node
        return None

    @staticmethod
    def _is_autocomplete_field(entry: Any) -> tuple[bool, bool]:
        """Detect combobox/autocomplete fields. Returns (is_combo, needs_js_wait).

        Mirrors browser-use tools/service.py:404-417. ``is_combo`` is True for any
        combobox-shaped field (drives the LLM hint); ``needs_js_wait`` is True only
        for the JS-driven subset (role=combobox or non-none aria-autocomplete) whose
        dropdowns populate asynchronously and need a ~0.4s settle before the next
        click — native <datalist> (list attr) and loose aria-haspopup render
        synchronously and are excluded from the wait.
        """
        attrs = getattr(entry, "attributes", {}) or {}
        if attrs.get("role") == "combobox":
            return True, True
        aria_ac = attrs.get("aria-autocomplete", "")
        if aria_ac and aria_ac != "none":
            return True, True
        if attrs.get("list"):
            return True, False  # native <datalist>: instant, no wait
        haspopup = attrs.get("aria-haspopup", "")
        if haspopup and haspopup != "false" and (attrs.get("aria-controls") or attrs.get("aria-owns")):
            return True, False
        return False, False

    async def _action_input_text(self, params: dict, browser: BrowserSession) -> ActionResult:
        # element_id is an alias for index (a backend node id from
        # find_elements(return_node_ids=True)). Exactly one required — guard
        # here since the registry does not validate the execute path.
        index = params.get("index")
        element_id = params.get("element_id")
        if (index is None) == (element_id is None):
            return ActionResult(
                error="input_text requires exactly one of `index` or `element_id`.",
            )
        params = {**params, "index": index if index is not None else element_id}
        entry, error = await self._get_element_by_index(params["index"], browser)
        if error:
            return error
        backend_id = entry.backend_node_id
        text = params["text"]
        clear = params.get("clear", True)

        # 1. Focus: highlight -> click_element, mapping the bool signal
        #    (mirrors _action_click — no silent success on focus failure).
        try:
            await browser.highlight_element(backend_id)
            clicked = await browser.click_element(backend_id)
        except Exception as e:
            return ActionResult(error=f"Input focus failed: {e}")
        if not clicked:
            return ActionResult(
                error=(
                    f"Could not focus element {params['index']} for input "
                    f"(no coordinates and JS click fallback failed; "
                    f"the element may be detached, hidden, or in a cross-origin iframe)"
                ),
            )
        await asyncio.sleep(0.1)  # retain the original focus settle

        # 2. Type: date/time/special inputs reject per-char key events, so
        #    assign directly via the native setter (_force_set_value) instead
        #    of _type_char. Mirrors browser-use _requires_direct_value_assignment
        #    / _set_value_directly.
        try:
            if _requires_direct_value_assignment(entry):
                if clear:
                    await browser._clear_text_field()
                await browser._force_set_value(text)
            else:
                await browser.type_text(text, clear=clear)
        except Exception as e:
            return ActionResult(error=f"Failed to type text into element {params['index']}: {e}")

        # 3. autocomplete/combobox: sleep ~0.4s on the JS-driven subset so the
        #    dropdown populates before the next action (browser-use service.py:812-819).
        is_combo, needs_js_wait = self._is_autocomplete_field(entry)
        if needs_js_wait:
            await asyncio.sleep(0.4)

        # 4. Value verification: read back activeElement and append a ⚠️ Note when
        #    it differs from the intended text (browser-use service.py:804-810).
        #    _read_active_text swallows exceptions and returns "" — safe to call.
        memory = self._describe_input(entry, params["index"], text)
        actual = await browser._read_active_text()
        if actual and actual != text:
            memory += (
                f"  ⚠️ Note: the field's actual value {actual!r} differs from "
                f"the intended {text!r}. The site may have reformatted, truncated, "
                f"or rejected the input — re-observe before continuing."
            )
        if is_combo:
            memory += (
                "  💡 autocomplete field — select from the JS-populated dropdown "
                "if applicable instead of typing the full value."
            )
        logger.info(memory)
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    async def _action_scroll(self, params: dict, browser: BrowserSession) -> ActionResult:
        direction = params.get("direction", "down")
        amount = int(params.get("amount", 3))
        try:
            # G5: scroll 现在返回 {vertical_percentage, at_edge}（G2 位置读取）
            position = await browser.scroll(direction, amount)
        except Exception as e:
            # G4: scroll 非幂等，CDP 失败=没滚，必须报 error（区别于 close_tab 的软成功）
            logger.warning("scroll(%s, %d) failed: %s", direction, amount, e)
            return ActionResult(error=f"Scroll failed: {e}")
        # G1 + G2: 回显方向/量 + 当前位置；已到边界则当轮提示
        memory = f"Scrolled {direction} {amount} viewport-heights"
        pct = position.get("vertical_percentage")
        if pct is not None:
            memory += f" ({pct}% down)"
        if position.get("at_edge"):
            memory += f" (already at {direction}, no further content)"
        logger.info(memory)
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    async def _action_search(self, params: dict, browser: BrowserSession) -> ActionResult:
        import urllib.parse

        query = params["query"]
        engine = params.get("engine", "baidu")
        encoded_query = urllib.parse.quote_plus(query)
        url = _SEARCH_ENGINE_URLS[engine].format(query=encoded_query)
        await browser.navigate(url)
        memory = f"Searched {engine.title()} for '{query}'"
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    async def _action_extract(self, params: dict, browser: BrowserSession) -> ActionResult:
        query = params["query"]
        extract_links = params.get("extract_links", True)
        extract_images = params.get("extract_images", True)
        start_from_char = params.get("start_from_char", 0)
        already_collected = params.get("already_collected")
        tr = self._truncation
        schema = getattr(self, "_extraction_schema", None)

        # 1) 源：CDP HTML → markdown（取代阶段一的 document.body.innerText）
        html_text = await browser.get_page_html(extract_links=extract_links, extract_images=extract_images)
        if not html_text:  # 降级到 execute_js outerHTML
            try:
                html_text = await browser.execute_js("document.documentElement.outerHTML") or ""
            except Exception as e:
                logger.warning("extract: HTML source failed: %s", e)
                html_text = ""
        if not html_text:
            return ActionResult(extracted_content="(empty page)")

        md = extract_clean_markdown(html_text, extract_links=extract_links, extract_images=extract_images)
        if not md.strip():
            return ActionResult(extracted_content="(empty page)")

        llm = getattr(self, "_extract_llm", None)
        if llm is None:
            # 未接 LLM（如脱离 Agent 直接用 Tools）——显式降级为截断 markdown 片段
            snippet = md[start_from_char:start_from_char + tr.extract_fallback_max_chars]
            return ActionResult(extracted_content=snippet or "(no content at offset)")

        # 2) 分页：取 start_from_char 所在的单块（一次 extract 只抽一块）
        chunks = chunk_markdown_by_structure(md, max_chars=tr.extract_chunk_max_chars)
        if start_from_char >= chunks[-1].end_index:
            return ActionResult(extracted_content="(no more content at this offset; extraction complete)")
        target_idx = 0
        for i, c in enumerate(chunks):
            if c.start_index <= start_from_char < c.end_index:
                target_idx = i
                break
        local_offset = max(0, start_from_char - chunks[target_idx].start_index)
        chunk_content = chunks[target_idx].content[local_offset:]

        # 3) 抽取（含去重 + 内层超时）
        try:
            result = await llm.extract(
                query,
                chunk_content,
                max_content_chars=tr.extract_chunk_max_chars,
                output_schema=schema,
                already_collected=already_collected,
                call_timeout=tr.extract_call_timeout or None,
            )
        except TimeoutError as e:
            logger.warning("extract: LLM call timed out: %s", e)
            return ActionResult(error=f"Extract timed out: {e}")
        except Exception as e:
            logger.warning("extract: LLM call failed: %s", e)
            return ActionResult(error=f"Extract failed: {e}")

        # 4) 分页进度（提示必须落在 500 字窗口内可见）
        next_offset = chunks[target_idx + 1].start_index if target_idx + 1 < len(chunks) else None
        remaining = sum(c.end_index - c.start_index for c in chunks[target_idx + 1:])
        hint = ""
        if next_offset is not None:
            hint = (f"[chunk {target_idx + 1}/{len(chunks)}; ~{remaining} chars remain; "
                    f"call extract again with start_from_char={next_offset} to continue]")

        # 5) 大结果分级落盘（仅按大小，与分页解耦；复用 write_file 直写姿势）
        saved_to = None
        if len(result) >= tr.extract_save_threshold:
            try:
                os.makedirs(tr.extract_output_dir, exist_ok=True)
                ext = "json" if schema else "md"
                fpath = os.path.join(tr.extract_output_dir, f"extract_{int(time.time() * 1000)}.{ext}")
                with open(fpath, "w", encoding="utf-8", newline="") as f:
                    f.write(result)
                saved_to = fpath
            except OSError as e:
                logger.warning("extract: save to file failed: %s", e)

        # 结构化结果保持 JSON 纯净；free-text 提示前置（防 __str__ 500 字截断）
        if schema is not None:
            if saved_to:
                visible = (f"Extraction ({len(result)} chars) saved to {saved_to}. "
                           f"Preview: {result[:200]}...\n{hint}").strip()
            else:
                visible = result  # 纯 JSON；hint 走 long_term_memory
        else:
            visible = (hint + "\n" + result) if hint else result
        mem_parts = []
        if saved_to:
            mem_parts.append(f"extract result saved: {saved_to}")
        if hint:
            mem_parts.append(hint)
        return ActionResult(extracted_content=visible, long_term_memory=" | ".join(mem_parts) or None)

    async def _action_send_keys(self, params: dict, browser: BrowserSession) -> ActionResult:
        keys = params["keys"]
        try:
            await browser.send_keys(keys)
        except Exception as e:
            # send_keys is NOT idempotent (a key can submit a form / trigger
            # navigation), so a CDP failure must surface as error — mirrors the
            # scroll pattern, unlike close_tab's soft-success degradation.
            logger.warning("send_keys(%r) failed: %s", keys, e)
            return ActionResult(error=f"Send keys failed: {e}")
        memory = f"Sent keys '{keys}'"
        logger.info(memory)
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    async def _action_switch_tab(self, params: dict, browser: BrowserSession) -> ActionResult:
        tab_id_suffix = params["tab_id"]
        tabs = await browser.get_tabs()
        matches = [t for t in tabs if t.target_id.endswith(tab_id_suffix)]
        if not matches:
            return ActionResult(
                error=f"No tab ending with '{tab_id_suffix}'. "
                      f"Open tabs: {self._summarize_tabs(tabs)}",
            )
        if len(matches) > 1:  # 后缀撞车：切错页风险，要求更长后缀/完整 target_id
            return ActionResult(
                error=f"Multiple tabs match '{tab_id_suffix}' ({len(matches)}). "
                      f"Use more characters or the full target_id. "
                      f"Matches: {self._summarize_tabs(matches)}",
            )
        target = matches[0]
        await browser.switch_tab(target.target_id)
        memory = f"Switched to tab [{tab_id_suffix}] {target.title} ({target.url})"
        logger.info(memory)
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    @staticmethod
    def _summarize_tabs(tabs: list) -> str:
        """把标签页列表摘要成短串，供 switch_tab 的 error/回显复用。"""
        items = []
        for t in tabs:
            title = (t.title or "").strip()[:40]
            url = (t.url or "").strip()[:60]
            items.append(f"[{t.target_id[-4:]}] {title} - {url}")
        return "; ".join(items)

    async def _action_close_tab(self, params: dict, browser: BrowserSession) -> ActionResult:
        tab_id_suffix = params.get("tab_id", "")
        tabs = await browser.get_tabs()  # G2: 轻量枚举，取代 get_state(include_screenshot=False)
        if tab_id_suffix:
            # 指定 tab_id：后缀匹配 + 冲突检测（G3）+ 未命中列出（G4）
            matches = [t for t in tabs if t.target_id.endswith(tab_id_suffix)]
            if not matches:
                return ActionResult(
                    error=f"No tab ending with '{tab_id_suffix}'. "
                          f"Open tabs: {self._summarize_tabs(tabs)}",
                )
            if len(matches) > 1:  # 后缀撞车：关错页风险，要求更长后缀/完整 target_id
                return ActionResult(
                    error=f"Multiple tabs match '{tab_id_suffix}' ({len(matches)}). "
                          f"Use more characters or the full target_id. "
                          f"Matches: {self._summarize_tabs(matches)}",
                )
            target = matches[0]
            target_id, id_echo, title, url = (
                target.target_id, tab_id_suffix, target.title, target.url,
            )
        else:
            # G6: 空 tab_id = 关当前页（保留原语义，补回显）
            if not browser.current_target_id:
                return ActionResult(error="No current tab to close")
            target_id = browser.current_target_id
            cur = next((t for t in tabs if t.target_id == target_id), None)
            id_echo = target_id[-4:]
            title = cur.title if cur else ""
            url = cur.url if cur else ""
        # 关闭（G1 回显 + G5 软降级，对齐 browser-use service.py:1011-1018）
        try:
            await browser.close_tab(target_id)
        except Exception as e:
            logger.warning("close_tab(%s) failed: %s", target_id, e)
            memory = f"Tab [{id_echo}] {title} ({url}) was already closed or invalid"
            return ActionResult(extracted_content=memory, long_term_memory=memory)
        memory = f"Closed tab [{id_echo}] {title} ({url})"
        logger.info(memory)
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    async def _action_wait(self, params: dict, browser: BrowserSession) -> ActionResult:
        await asyncio.sleep(float(params.get("seconds", 3)))
        return ActionResult()

    async def _action_go_back(self, params: dict, browser: BrowserSession) -> ActionResult:
        try:
            target_url = await browser.go_back()
        except Exception as e:
            return ActionResult(error=f"Failed to go back: {e}")

        if target_url is None:
            # 无历史可退（currentIndex <= 0）——明确告知，避免 LLM 误以为已后退
            return ActionResult(error="No previous page in history to go back to")

        # 轻量健康检查：SPA 后退未渲染给一次重试机会（仍空仅 warning，不硬失败）
        await self._go_back_health_check(target_url, browser)

        memory = f"Navigated back to {target_url}"
        logger.info(memory)
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    async def _go_back_health_check(self, target_url: str | None, browser: BrowserSession) -> None:
        """后退后轻量空 DOM 检测（用户选：轻量，不 reload、不硬失败）。

        与 _navigate_health_check 的差异：go_back 没有"reload 同 URL"的干净机制
        （重新 navigate 会新增历史项），故仅做一次重试等待，持续空只 warning 不抛错，
        交由 LLM 下一轮 get_state 自行感知。仅当后退目标为 http(s) 时触发。
        """
        state = await browser.get_state(include_screenshot=False)
        url_is_http = state.url.lower().startswith(("http://", "https://"))
        if not (url_is_http and self._dom_appears_empty(state)):
            return

        logger.warning(
            "Empty DOM after going back to %s, waiting %.0fs and rechecking",
            target_url, _NAVIGATE_EMPTY_RETRY_WAIT,
        )
        await asyncio.sleep(_NAVIGATE_EMPTY_RETRY_WAIT)
        state = await browser.get_state(include_screenshot=False)
        if state.url.lower().startswith(("http://", "https://")) and self._dom_appears_empty(state):
            logger.warning(
                "Still empty after going back to %s; SPA may still be rendering. "
                "Not failing hard (no clean reload for history navigation).",
                target_url,
            )

    async def _action_find_elements(self, params: dict, browser: BrowserSession) -> ActionResult:
        selector = params["selector"]
        max_results = params.get("max_results", 50)
        offset = params.get("offset", 0)
        return_node_ids = params.get("return_node_ids", False)
        try:
            if return_node_ids:
                # Phase 2 (item ②): resolve to stable backendNodeIds via
                # DOM.performSearch so results can feed click/input_text
                # directly (index===backend_id). Heavier; no text returned.
                data = await browser.find_elements_node_ids(
                    selector, max_results=max_results, offset=offset,
                )
                formatter = _format_node_id_results
            else:
                data = await browser.find_elements(
                    selector,
                    attributes=params.get("attributes"),
                    max_results=max_results,
                    offset=offset,
                    include_text=params.get("include_text", True),
                    first_only=params.get("first_only", False),
                    include_geometry=params.get("include_geometry", False),
                )
                formatter = _format_find_results
        except Exception as e:
            # CDP layer failure (connection drop / invalid selector surfaced as
            # {error} -> RuntimeError) = tool execution failure; surface a
            # find_elements-specific error rather than the generic Tools.execute
            # fallback. Aligns with _action_find_text / _action_search_page.
            logger.warning("find_elements(%r) failed: %s", selector, e)
            return ActionResult(error=f"Find elements failed: {e}")
        total = data.get("total", 0)
        if total == 0:
            # Soft echo: "selector matched nothing" is actionable info (LLM can
            # fix the selector / wait / accept absence), not a tool failure —
            # aligns with find_text / search_page / browser-use.
            msg = f'No elements found matching "{selector}"'
            logger.info(msg)
            return ActionResult(extracted_content=msg, long_term_memory=msg)
        formatted = formatter(data, selector)
        # 大结果分级落盘（镜像 _action_search_page / _action_extract；OSError 不失败只 warning）
        tr = self._truncation
        saved_to = None
        if len(formatted) >= tr.find_elements_save_threshold:
            try:
                os.makedirs(tr.find_elements_output_dir, exist_ok=True)
                fpath = os.path.join(tr.find_elements_output_dir, f"find_elements_{int(time.time() * 1000)}.txt")
                with open(fpath, "w", encoding="utf-8", newline="") as f:
                    f.write(formatted)
                saved_to = fpath
            except OSError as e:
                logger.warning("find_elements: save to file failed: %s", e)
        if saved_to:
            visible = (f"Find results ({len(formatted)} chars) saved to {saved_to}. "
                       f"Preview: {formatted[:200]}...").strip()
        else:
            visible = formatted
        nid_suffix = " (node ids)" if return_node_ids else ""
        memory = f'Found {total} element{"s" if total != 1 else ""} matching "{selector}"{nid_suffix}.'
        if saved_to:
            memory += f" Results saved: {saved_to}"
        logger.info(memory)
        return ActionResult(extracted_content=visible, long_term_memory=memory)

    async def _action_find_text(self, params: dict, browser: BrowserSession) -> ActionResult:
        text = params["text"]
        nth = params.get("nth", 1)
        case_sensitive = params.get("case_sensitive", False)
        highlight = params.get("highlight", "box")
        try:
            info = await browser.find_text(
                text, nth=nth, case_sensitive=case_sensitive, highlight=highlight,
            )
        except Exception as e:
            # CDP layer failure (connection drop / DOM domain error) = tool
            # execution failure; surface a find_text-specific error rather
            # than the generic Tools.execute fallback.
            logger.warning("find_text(%r, nth=%d) failed: %s", text, nth, e)
            return ActionResult(error=f"Find text failed: {e}")
        if not info.get("found"):
            if info.get("reason") == "nth_exceeds":
                # Text is on the page but nth exceeds the visible match count:
                # actionable (caller can lower nth), not a tool failure.
                msg = (
                    f"Text '{text}' found but only {info.get('visible_total')} visible "
                    f"match(es) ({info.get('total')} total via {info.get('method')}) "
                    f"— asked for match {info.get('requested_nth')}, try a smaller nth"
                )
                logger.info(msg)
                return ActionResult(extracted_content=msg, long_term_memory=msg)
            # Soft echo: "text not on the page" is actionable info (the LLM
            # can scroll / switch tab / accept the text is absent), not a tool
            # failure — aligns with browser-use and the sibling search_page
            # action (both return extracted_content, not error, on a miss).
            msg = f"Text '{text}' not found on page"
            logger.info(msg)
            return ActionResult(extracted_content=msg, long_term_memory=msg)
        method = info.get("method")
        tag = info.get("tag")
        total = info.get("total")
        if total and total > 1:
            # Multi-match: report which match out of how many (G8).
            counts = f"match {info.get('match_index')} of {info.get('visible_total')} visible, {total} total"
            if tag:
                memory = f"Scrolled to text '{text}' into view ({counts}, found in <{tag}>, via {method})"
            else:
                memory = f"Scrolled to text '{text}' into view ({counts}, via {method})"
        else:
            # Single match (total==1) or JS fallback (no count): P0-style echo.
            if tag:
                memory = f"Scrolled to text '{text}' into view (found in <{tag}>, via {method})"
            else:
                memory = f"Scrolled to text '{text}' into view (via {method})"
        info_highlight = info.get("highlight")
        if info_highlight and info_highlight != "box":
            memory += f" ({info_highlight} highlight)"
        logger.info(memory)
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    async def _action_screenshot(self, params: dict, browser: BrowserSession) -> ActionResult:
        fmt: str = params.get("format", "png")
        quality = params.get("quality")
        clip = params.get("clip")
        full_page: bool = params.get("full_page", False)
        save_path: str = params.get("save_path", "")

        try:
            screenshot_bytes = await browser.take_screenshot(
                format=fmt, quality=quality, clip=clip,
                full_page=full_page, wait_settle=full_page,
            )
        except Exception as e:
            logger.warning("screenshot action failed: %s", e)
            return ActionResult(error=f"Screenshot failed: {e}")

        if save_path:
            try:
                with open(save_path, "wb") as f:
                    f.write(screenshot_bytes)
            except OSError as e:
                return ActionResult(error=f"Failed to save screenshot to {save_path}: {e}")
            return ActionResult(extracted_content=f"Screenshot saved to {save_path} ({len(screenshot_bytes)} bytes)")

        meta = f"format={fmt}, {len(screenshot_bytes)} bytes"
        if full_page:
            meta += ", full_page"
        if clip:
            meta += f", clip={clip.get('width')}x{clip.get('height')}"
        return ActionResult(extracted_content=f"Screenshot captured ({meta}) but not saved (no save_path).")

    async def _action_save_as_pdf(self, params: dict, browser: BrowserSession) -> ActionResult:
        path: str = params["path"]
        paper_format: str = params.get("paper_format", "letter")
        landscape: bool = params.get("landscape", False)
        print_background: bool = params.get("print_background", True)
        scale: float = params.get("scale", 1.0)

        try:
            pdf_bytes = await browser.print_to_pdf(
                paper_format=paper_format,
                landscape=landscape,
                print_background=print_background,
                scale=scale,
            )
        except Exception as e:
            logger.warning("save_as_pdf action failed: %s", e)
            return ActionResult(error=f"Failed to generate PDF: {e}")

        try:
            os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
            with open(path, "wb") as f:
                f.write(pdf_bytes)
        except OSError as e:
            return ActionResult(error=f"Failed to save PDF to {path}: {e}")

        meta = f"paper={paper_format}, {len(pdf_bytes)} bytes"
        if landscape:
            meta += ", landscape"
        return ActionResult(extracted_content=f"PDF saved to {path} ({meta})")

    async def _action_dropdown_options(self, params: dict, browser: BrowserSession) -> ActionResult:
        """读取指定 index 的下拉元素全部 option（native <select> / ARIA menu·listbox
        / custom class / combobox / 子树搜索）。

        action 层做廉价预分类：native <select> 沿用 P0 的 fetch_select_options（零改动）；
        combobox（_is_autocomplete_field + aria-controls/owns）走 session 的 Python flow
        expand_and_fetch_combobox_options（展开→读→收起）；其余委托 session 轻量 dispatcher
        fetch_dropdown_options（顺序试 aria→custom→子树）。所有类型共用 _format_options_result
        的 json 编码 + 序号 + 用法提示 + 短/长回显；source 折进 long_term_memory 作诊断通道。
        对齐 _describe_click / _describe_upload 的成功回显 + try/except 软降级规范。
        """
        index = params["index"]
        entry, error = await self._get_element_by_index(index, browser)
        if error:
            return error

        tag = (getattr(entry, "tag_name", "") or "").upper()
        backend_id = getattr(entry, "backend_node_id", None)
        is_combo, _ = self._is_autocomplete_field(entry)
        attrs = getattr(entry, "attributes", {}) or {}

        try:
            # native <select>：P0 路径零改动（source=native 无 via 后缀，与 P0 字节一致）
            if tag == "SELECT":
                raw_options = await browser.fetch_select_options(backend_id)
                return self._format_options_result(raw_options, entry, index, "native")
            # combobox（aria-controls 独立 listbox）：Python flow（展开→读→收起）
            if is_combo and (attrs.get("aria-controls") or attrs.get("aria-owns")):
                raw_options = await browser.expand_and_fetch_combobox_options(backend_id)
                return self._format_options_result(raw_options, entry, index, "combobox")
            # 其余：session 层 dispatcher（aria / custom / 子树）
            dispatched = await browser.fetch_dropdown_options(backend_id)
            if dispatched["source"] is None:
                # 真阴性：非任何已知下拉类型，回退到 P0 风格的友好 error 提示手动展开
                return ActionResult(
                    error=(
                        f"Index {index} is a [{tag}] element, not a recognized dropdown "
                        f"(native <select>, ARIA listbox/menu, custom dropdown, or combobox). "
                        f"Use click to expand and read options manually."
                    ),
                )
            return self._format_options_result(
                dispatched["options"], entry, index, dispatched["source"],
            )
        except Exception as e:
            return ActionResult(error=f"Failed to read dropdown options: {e}")

    async def _action_select_dropdown(self, params: dict, browser: BrowserSession) -> ActionResult:
        """Select an option in the dropdown at the given index.

        P0 scoped the native <select> chain to the target's backend_node_id (DOM.resolveNode
        + Runtime.callFunctionOn), fixing the global querySelectorAll('select')[0] bug and
        porting browser-use's native selection chain (focus -> set value three ways ->
        input/change/blur -> readback verify -> click fallback on framework reversion). P1
        widens this to multi-type dispatch (mirrors _action_dropdown_options): native
        <select> -> set_select_option; combobox (aria-controls) -> set_combobox_option
        (Python flow); otherwise the session write-dispatcher set_dropdown_option classifies
        (aria/custom/subtree, reusing read-side fetch_dropdown_options) and routes to the
        matching setter. On a miss, echoes the available options back for the LLM to
        self-correct. Mirrors the _describe_dropdown + try/except soft-degrade convention of
        dropdown_options. The success/miss/error three-segment handling is shared by all
        setter types (D2 — uniform dict shape, zero per-type branching here).
        """
        index = params["index"]
        entry, error = await self._get_element_by_index(index, browser)
        if error:
            return error

        # Multi-type dispatch (P1, mirrors _action_dropdown_options). native <select> ->
        # set_select_option (P0 path, zero change); combobox (aria-controls) -> set_combobox_option
        # (Python flow); otherwise the session write-dispatcher set_dropdown_option classifies
        # (aria/custom/subtree, reusing read-side fetch_dropdown_options) and routes to the
        # matching setter. Tag guard widened from P0's "tag != SELECT hard-reject".
        tag = (getattr(entry, "tag_name", "") or "").upper()
        backend_id = getattr(entry, "backend_node_id", None)
        value = params["value"]
        is_combo, _ = self._is_autocomplete_field(entry)
        attrs = getattr(entry, "attributes", {}) or {}

        try:
            if tag == "SELECT":
                # native <select>：P0 路径零改动
                result = await browser.set_select_option(backend_id, value)
            elif is_combo and (attrs.get("aria-controls") or attrs.get("aria-owns")):
                # combobox（aria-controls 独立 listbox）：Python flow（展开→定位→写→收起）
                result = await browser.set_combobox_option(backend_id, value)
            else:
                # aria / custom / 子树：session 写 dispatcher（复用读侧分类，读写零漂移）
                result = await browser.set_dropdown_option(backend_id, value)
                if result.get("source") is None:
                    # 真阴性：非任何已知下拉类型（与 dropdown_options 真阴性一致）
                    return ActionResult(
                        error=(
                            f"Index {index} is a [{tag}] element, not a recognized dropdown "
                            f"(native <select>, ARIA listbox/menu, custom dropdown, or combobox). "
                            f"Use dropdown_options to list available options first."
                        ),
                    )
        except Exception as e:
            return ActionResult(error=f"Failed to select option: {e}")

        # Success echo (short/long split, mirrors _describe_dropdown + json.dumps).
        desc = self._describe_dropdown(entry, index)
        if result.get("success"):
            message = result.get("message", f"Selected option: {value}")
            memory = f"Selected {json.dumps(value)} in {desc}"
            return ActionResult(extracted_content=message, long_term_memory=memory)

        # Miss / framework block (and click fallback also failed): echo the
        # available options so the LLM can retry with a correct value.
        available = result.get("availableOptions") or []
        if available:
            lines = [
                f"{i}: text={json.dumps(o.get('text', ''))}, value={json.dumps(o.get('value', ''))}"
                for i, o in enumerate(available)
            ]
            extracted = "\n".join(lines) + "\n" + f"Use the value in select_dropdown(index={index}, value=...)"
            memory = f"Couldn't select {json.dumps(value)} in {desc} (not an available option)"
            return ActionResult(extracted_content=extracted, long_term_memory=memory)

        # Bare failure with no available options (e.g. a JS-returned pure error).
        err = result.get("error", f"Failed to select option: {value}")
        return ActionResult(error=err)

    async def _action_upload_file(self, params: dict, browser: BrowserSession) -> ActionResult:
        file_path = params["path"]

        # 1. 路径白名单 + 文件存在/非空校验（保持原逻辑）
        allowed = self._allowed_upload_paths
        if allowed:
            if not any(file_path.startswith(p) for p in allowed):
                return ActionResult(
                    error=f"File path not in allowed upload paths: {file_path}",
                )
        if not os.path.isfile(file_path):
            return ActionResult(error=f"File not found: {file_path}")
        if os.path.getsize(file_path) == 0:
            return ActionResult(error=f"File is empty: {file_path}")

        # 2. 元素查找（保持原逻辑）
        entry, error = await self._get_element_by_index(params["index"], browser)
        if error:
            return error

        # 3. 判定目标是否本身 file input；确定正确的 file input
        tag = entry.tag_name.upper()
        attrs = entry.attributes
        is_file_input = tag == "INPUT" and attrs.get("type", "").lower() == "file"

        backend_id = entry.backend_node_id
        file_input_ids: list[int] = []
        file_inputs_meta: list = []
        if self._cached_browser_state and self._cached_browser_state.dom_state:
            file_input_ids = list(
                self._cached_browser_state.dom_state.file_input_backend_ids,
            )
            file_inputs_meta = list(
                self._cached_browser_state.dom_state.file_inputs_meta,
            )
        upload_note = ""

        if is_file_input and len(file_input_ids) > 1:
            # 多个 file input 共存（抖音式封面编辑器：除真实封面 input 外还有"收藏
            # 封面"等无关 input）。但仍信任 agent 指定的 index 直接 setFileInputFiles——
            # 抖音封面无 <label>（issue #34），"改点可见上传按钮"这条路走不通，硬拒绝
            # 反而导致 0% 成功率（实测 master 直接 set 能传）。改为软警告：仍上传到
            # agent 指定的 index，同时点名「可见 + upload 容器内」的候选 input；
            # 若本次未生效（页面无变化 = 命中隐藏诱饵，或弹出收藏框）就改试这些候选。
            live_candidates = [
                fi.backend_node_id for fi in file_inputs_meta
                if getattr(fi, "visible", False) and getattr(fi, "upload_ancestor", False)
            ]
            cand_hint = (
                f" Likely-live candidates (visible + upload container): {live_candidates}."
                if live_candidates else ""
            )
            upload_note = (
                f"  ⚠️ Page has {len(file_input_ids)} file inputs; uploaded to the one "
                f"you specified (index {params['index']}).{cand_hint} If the site reacted "
                f"wrongly (a 收藏封面/favorite-cover modal popped, or nothing changed = "
                f"you hit a hidden decoy input), retry upload_file on the correct visible "
                f"upload area."
            )

        if not is_file_input:
            if not file_input_ids:
                return ActionResult(
                    error="Element is not a file input and no file input found on page",
                )
            if len(file_input_ids) == 1:
                # Exactly one file input on the page — unambiguous, use it directly.
                backend_id = file_input_ids[0]
                upload_note = (
                    f"  ℹ️ index {params['index']} is not a file input; uploaded to the "
                    f"only file input on the page (backendNodeId={backend_id})."
                )
            else:
                # 多个 file input 且目标是 dropzone/按钮：点其附近的 <label> 上传触发器
                # 而非目标本身。<label for=input> 是浏览器原生行为，保留 user-gesture，
                # 能触发 Page.fileChooserOpened；而 drag-area/div/button 的 JS
                # input.click() 会丢失 gesture、不触发 chooser（issue #34 抖音封面）。
                # 命中 chooser 的瞬态 input 才是页面真正为这次上传关联的 input。
                label_bid = _find_upload_label_near(entry) or entry.backend_node_id
                discovered = await browser.discover_file_input_via_click(label_bid)
                if discovered is not None:
                    backend_id = discovered
                    upload_note = (
                        f"  ℹ️ index {params['index']} is not a file input; clicked its "
                        f"upload button and uploaded to the file input the page opened "
                        f"(backendNodeId={backend_id})."
                    )
                else:
                    return ActionResult(error=(
                        f"Element {params['index']} is not a file input, and clicking its "
                        f"upload button did not open a file chooser — it likely uses a "
                        f"custom upload dialog. Open that dialog and click its 上传图片/"
                        f"选择文件 button first, then call upload_file again on that button "
                        f"so the correct file input is used."
                    ))
        # else: is_file_input（无论唯一还是多 input）→ backend_id 已是 entry.backend_node_id，直传（多 input 时上面已附软警告）

        logger.info(
            "upload_file: index=%d, tag=%s, type=%s, backend_node_id=%s, "
            "is_file_input=%s, resolved_backend_id=%s, available_file_inputs=%s",
            params["index"], tag, attrs.get("type", ""), entry.backend_node_id,
            is_file_input, backend_id, file_input_ids,
        )

        # 4. 高亮 + 上传（对齐 _action_click：共用 try，统一映射；highlight best-effort）
        try:
            await browser.highlight_element(backend_id)
            await browser.set_file_input(
                backend_node_id=backend_id,
                file_path=file_path,
                file_input_backend_ids=file_input_ids if not is_file_input else None,
            )
        except Exception as e:
            return ActionResult(error=f"File upload failed: {e}")

        # 5. 成功回显 + 目标来源说明（直选/页面选中）+ accept 软校验
        memory = self._describe_upload(entry, params["index"], file_path)
        if upload_note:
            memory += upload_note

        if is_file_input:
            file_input_entry = entry
        else:
            file_input_entry = self._find_node_by_backend_id(
                backend_id,
                self._cached_browser_state.dom_state if self._cached_browser_state else None,
            )
        accept_attr = None
        if file_input_entry is not None:
            accept_attr = (getattr(file_input_entry, "attributes", {}) or {}).get("accept")
        if accept_attr and not _file_matches_accept(file_path, accept_attr):
            memory += (
                f"  ℹ️ Note: file extension does not match this input's "
                f"accept={accept_attr!r}. The file was uploaded successfully regardless "
                f"— browsers do not enforce accept (it is advisory only). No retry needed."
            )

        logger.info(memory)
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    async def _action_write_file(self, params: dict, browser: BrowserSession) -> ActionResult:
        path = params["path"]
        content = params["content"]
        append = params.get("append", False)
        trailing_newline = params.get("trailing_newline", True)
        leading_newline = params.get("leading_newline", False)
        # 阶段二（二.B / 二.D）：编码与行尾翻译参数（默认复现阶段一行为）。
        enc = params.get("encoding") or "utf-8"
        newline = params.get("newline", "")

        # 阶段二（二.C）：写路径白名单（镜像 allowed_upload_paths 前缀匹配）。
        # None = 不启用、全放行；置于 makedirs/open 之前 fail fast，避免在 jail 外建目录/tmp。
        allowed = self._allowed_write_paths
        if allowed and not any(path.startswith(p) for p in allowed):
            return ActionResult(error=f"File path not in allowed write paths: {path}")

        # 换行簿记在 action 层（对齐 browser-use service.py:1691-1694），但 trailing
        # 采用守卫式（幂等、不双换行、不破坏 CRLF —— "foo\r\n".endswith("\n") 为 True）。
        if leading_newline:
            content = "\n" + content
        if trailing_newline and not content.endswith("\n"):
            content = content + "\n"

        # 阶段二（二.A）：overwrite 走原子写（tmp + os.replace，同目录→同卷原子），
        # 进程崩溃不留半个文件；append 刻意保持 open(path,"a") 直接追加（O(1)，非崩溃安全）。
        tmp = path + ".tmp"
        try:
            os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
            if append:
                # newline= 关闭文本模式换行翻译（见 2.D）；append 非原子。
                with open(path, "a", encoding=enc, newline=newline) as f:
                    f.write(content)
            else:
                with open(tmp, "w", encoding=enc, newline=newline) as f:
                    f.write(content)
                os.replace(tmp, path)
        except LookupError as e:
            # 非法 encoding 名（非 OSError 子类）：单独兜底，避免冒泡到通用 catch。
            # open() 可能已创建 tmp（overwrite）/path（append）残留，清理之。
            if not append and os.path.exists(tmp):
                try:
                    os.remove(tmp)
                except OSError:
                    pass
            logger.warning("write_file(%r) unknown encoding %r: %s", path, enc, e)
            return ActionResult(error=f"Unknown encoding {enc!r}: {e}")
        except OSError as e:
            # 仅 overwrite 路径会留下 tmp 残骸；清理之（吞二次 OSError）。分级错误（对齐
            # save_as_pdf:980-985），不冒泡到 Tools.execute 通用 catch（actions.py:260-262）。
            if not append and os.path.exists(tmp):
                try:
                    os.remove(tmp)
                except OSError:
                    pass
            logger.warning("write_file(%r) failed: %s", path, e)
            return ActionResult(error=f"Failed to write file {path}: {e}")

        written = len(content.encode(enc))
        action_word = "Appended" if append else "Wrote"
        memory = f"{action_word} {written} bytes to {path}"
        if enc != "utf-8":
            memory += f" (encoding: {enc})"
        logger.info(memory)
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    async def _action_read_file(self, params: dict, browser: BrowserSession) -> ActionResult:
        path = params["path"]
        # 阶段二（二.C）：读路径白名单（镜像 allowed_write_paths；None=全放行）。
        # 置于 sniff/open 之前 fail fast，与 write_file/replace_file 对称。
        allowed = self._allowed_read_paths
        if allowed and not any(path.startswith(p) for p in allowed):
            return ActionResult(error=f"File path not in allowed read paths: {path}")

        # 阶段二（二.B）：二进制嗅探（在文本解码前）。命中富文档转 二.D，
        # 命中不支持二进制给可操作 error（早拒，优于 UnicodeDecodeError 啰嗦堆栈）。
        try:
            kind = _sniff_file_kind(path)
        except FileNotFoundError:
            return ActionResult(error=f"File not found: {path}")
        except OSError as e:
            # path 指向目录(IsADirectoryError/PermissionError) 等 → 嗅探的 open 即抛。
            logger.warning("read_file(%r) sniff failed: %s", path, e)
            return ActionResult(error=f"Failed to read file {path}: {e}")
        if kind == "binary":
            logger.warning("read_file(%r) rejected binary kind", path)
            return ActionResult(error=f"{path} looks like a binary file; read_file reads UTF-8 text, "
                                     "PDF, DOCX, or images (PNG/JPEG/GIF/WebP).")
        if kind in ("pdf", "docx", "image"):
            return await self._read_rich_document(path, kind, params)

        # kind == "text" → 文本解码（newline/encoding 阶段一行为）。
        enc = params.get("encoding") or "utf-8"
        newline = params.get("newline", "")
        try:
            # newline="" 关闭 universal-newline 翻译（对齐 replace_file / write_file）：
            # 读时保留原始 \r\n，避免 CRLF 被压成 \n；行尾字节级不变，便于后续 replace_file
            # 用含 \r\n 的 old 精确匹配。newline="\n"/None 则启用 universal-newline。
            with open(path, "r", encoding=enc, newline=newline) as f:
                content = f.read()
        except FileNotFoundError:
            return ActionResult(error=f"File not found: {path}")
        except UnicodeDecodeError as e:
            # 读取侧特有（对齐 replace_file）：文件不是合法 enc 编码。
            logger.warning("read_file(%r) decode failed: %s", path, e)
            return ActionResult(error=f"Failed to decode {path} as {enc}: {e}")
        except LookupError as e:
            # 非法 encoding 名（非 OSError 子类）：单独兜底。
            logger.warning("read_file(%r) unknown encoding %r: %s", path, enc, e)
            return ActionResult(error=f"Unknown encoding {enc!r}: {e}")
        except OSError as e:
            # 分级错误（对齐 replace_file）：磁盘/只读/锁定 → 明确 error + warning，
            # 不冒泡到 Tools.execute 通用 catch（actions.py:260-262）。
            logger.warning("read_file(%r) failed: %s", path, e)
            return ActionResult(error=f"Failed to read file {path}: {e}")

        total_bytes = len(content.encode(enc))
        return self._window_and_echo(content, path, params, total_bytes)

    def _window_and_echo(self, content: str, path: str, params: dict, total_bytes: int) -> ActionResult:
        """阶段二（二.A）：offset/limit 字符级分页 + 截断 footer + 字符/字节回显。

        供文本路径与 二.D 的 PDF/DOCX 路径复用（动作体只编排，对齐 browser-use
        「service.py 只编排、逻辑下沉」思想）。字符级 offset 与 read_file_max_chars
        及阶段一 footer「showing X of Y chars」一致，使 offset=read_file_max_chars
        成为截断后的续读起点。
        """
        total_chars = len(content)
        if not content:
            # 软提示（对齐 replace_file soft-miss / search_page）：空内容不是错误，
            # 但要让 LLM 知道"读到空"，而非 __str__ 兜底成 "OK"。
            msg = f"{path} is empty (0 bytes)"
            logger.info(msg)
            return ActionResult(extracted_content=msg, long_term_memory=msg)
        offset = params.get("offset", 0)
        limit = params.get("limit")
        max_chars = self._truncation.read_file_max_chars
        window = limit if (limit is not None and limit < max_chars) else max_chars
        if offset >= total_chars:
            # offset 越过文件尾：软提示（非 error），对齐 empty soft-miss。
            msg = f"offset {offset} is at or past end of {path} ({total_chars} chars); nothing to read"
            logger.info(msg)
            return ActionResult(extracted_content=msg, long_term_memory=msg)
        content_window = content[offset:offset + window]
        shown = len(content_window)
        remaining = total_chars - offset - shown
        if remaining > 0:
            end = offset + shown
            extracted = (
                content_window
                + f"\n[...truncated: showing {shown} of {total_chars} chars "
                f"from offset {offset} ({total_bytes} bytes total); "
                f"use offset={end} to continue]"
            )
            memory = (
                f"Read {path} ({shown} of {total_chars} chars from offset {offset}, "
                f"{total_bytes} bytes; truncated)"
            )
        else:
            extracted = content_window
            if offset > 0:
                memory = (
                    f"Read {path} chars {offset}-{offset + shown} of {total_chars} "
                    f"({total_bytes} bytes; final page)"
                )
            else:
                memory = f"Read {path} ({total_chars} chars, {total_bytes} bytes)"
        logger.info(memory)
        return ActionResult(extracted_content=extracted, long_term_memory=memory)

    async def _read_rich_document(self, path: str, kind: str, params: dict) -> ActionResult:
        """阶段二（二.D）：富文档分派（由 _action_read_file 的二进制嗅探路由而来）。

        - PDF (pypdf) / DOCX (python-docx)：抽文本走 _window_and_echo（与文本路径
          同款 offset/limit 分页）；依赖做成 optional extra ``[docs]``，缺则给安装提示。
        - image：metadata['images'] 当前是死代码（evaluate 写了但 agent loop 从不读，
          LLM 只收 extracted_content 文本），故**暂不接线 vision 通道**，只返回可操作
          提示（不堆无用 base64）。待 agent loop 统一接线 images 通道后再启用。
        """
        if kind == "image":
            mime = mimetypes.guess_type(path)[0] or "application/octet-stream"
            try:
                nbytes = os.path.getsize(path)
            except OSError:
                nbytes = 0
            msg = (
                f"{path} is an image ({mime}, {nbytes} bytes). read_file cannot inline images "
                "yet (vision channel not wired); re-save as text/PDF, or use a vision-capable flow."
            )
            logger.info("read_file(%r) image detected; vision not wired", path)
            return ActionResult(extracted_content=msg, long_term_memory=msg)
        if kind == "pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                return ActionResult(
                    error=f"{path} is a PDF; install the 'docs' extra to enable PDF reading: "
                    "uv pip install -e .[docs]"
                )
            try:
                pages = PdfReader(path).pages
                parts = [
                    f"--- page {i + 1}/{len(pages)} ---\n{(p.extract_text() or '')}"
                    for i, p in enumerate(pages)
                ]
                text = "\n\n".join(parts)
            except Exception as e:  # pypdf 抛多种异常（PdfReadError 等），统一降级
                logger.warning("read_file(%r) pdf parse failed: %s", path, e)
                return ActionResult(error=f"Failed to parse PDF {path}: {e}")
        else:  # docx
            try:
                from docx import Document
            except ImportError:
                return ActionResult(
                    error=f"{path} is a DOCX; install the 'docs' extra to enable DOCX reading: "
                    "uv pip install -e .[docs]"
                )
            try:
                doc = Document(path)
                text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            except Exception as e:  # python-docx 抛多种异常，统一降级
                logger.warning("read_file(%r) docx parse failed: %s", path, e)
                return ActionResult(error=f"Failed to parse DOCX {path}: {e}")
        return self._window_and_echo(text, path, params, len(text.encode("utf-8")))

    async def _action_replace_file(self, params: dict, browser: BrowserSession) -> ActionResult:
        path = params["path"]
        old = params["old"]
        new = params["new"]
        if not old:
            # min_length=1 覆盖 schema + 直接构造；运行时 registry 不校验 params，
            # 此守卫兜底 execute 路径，避免 str.replace("", x) 在每字符间插入而膨胀文件。
            return ActionResult(error="replace_file 'old' must be a non-empty string")

        # 阶段二参数（registry 不校验 execute 路径，params 是 raw dict，运行时守卫）。
        enc = params.get("encoding") or "utf-8"
        newline = params.get("newline", "")
        regex = params.get("regex", False)
        case_sensitive = params.get("case_sensitive", True)  # 默认 True（分叉 search_page）
        count = params.get("count", None)
        expected_count = params.get("expected_count", None)
        backup = params.get("backup", False)

        # 运行时守卫：count 必须为 None 或正整数；expected_count 必须为 None 或非负整数。
        # bool 是 int 子类——True/False 须显式拒，否则被当 1/0。
        if count is not None and (not isinstance(count, int) or isinstance(count, bool) or count < 1):
            return ActionResult(error=f"replace_file 'count' must be a positive integer (got {count!r})")
        if expected_count is not None and (not isinstance(expected_count, int) or isinstance(expected_count, bool) or expected_count < 0):
            return ActionResult(error=f"replace_file 'expected_count' must be a non-negative integer (got {expected_count!r})")

        # 阶段二（二.C）：写路径白名单（replace_file 也是写）。.bak/.tmp 与 path 同目录同前缀，
        # path 放行即隐含 bak/tmp 放行，无需单独检查。
        allowed = self._allowed_write_paths
        if allowed and not any(path.startswith(p) for p in allowed):
            return ActionResult(error=f"File path not in allowed write paths: {path}")

        # 仅在 regex 或大小写不敏感时走 re；默认（literal+case_sensitive）保留阶段一
        # str.count/str.replace 路径，字节级不变、零回归。
        use_re = bool(regex) or not case_sensitive

        def _literal_replacer(literal: str):
            # 大小写不敏感但非 regex：把 new 当不透明字面量，不展开 \1 / \g<name>。
            return lambda m: literal

        # 阶段二（二.A）：原子写 tmp（同目录→同卷原子），写段用 tmp + os.replace。
        tmp = path + ".tmp"
        bak = path + ".bak"
        try:
            # newline="" 关闭 universal-newline 翻译（对齐 write_file:1263）：
            # 读时保留原始 \r\n，写时 \n 不被译成 \r\n。原 LF 保持 LF、原 CRLF 保持
            # CRLF，行尾字节级不变；含 \r\n 字面量的 old/new 也不会被压缩成 \n。
            with open(path, "r", encoding=enc, newline=newline) as f:
                content = f.read()

            # 计算原始匹配总数（expected_count 基准，也是软失败判定）。
            if use_re:
                flags = 0 if case_sensitive else re.IGNORECASE
                try:
                    pattern = re.compile(old if regex else re.escape(old), flags)
                except re.error as e:
                    logger.warning("replace_file(%r) invalid regex %r: %s", path, old, e)
                    return ActionResult(error=f"Invalid regex pattern {old!r}: {e}")
                raw_total = len(pattern.findall(content))
            else:
                raw_total = content.count(old)

            # expected_count：写前守卫，比对原始总数；不匹配则文件不动（typo-guard）。
            if expected_count is not None and raw_total != expected_count:
                msg = (f"replace_file expected {expected_count} match(es) for {old!r} in {path}, "
                       f"found {raw_total}; file unchanged")
                logger.info(msg)
                return ActionResult(error=msg)

            # 软失败：raw_total==0（含 expected_count=0 & actual=0）→ 不写、成功语义。
            # 修正 browser-use 的静默"Successfully replaced"缺陷——绝不假装改了。
            if raw_total == 0:
                msg = f"No occurrences of {old!r} found in {path}; file unchanged"
                logger.info(msg)
                return ActionResult(extracted_content=msg, long_term_memory=msg)

            # 备份：读成功、校验通过后再复制原始内容到 .bak。失败致命——不写。
            if backup:
                try:
                    shutil.copy2(path, bak)
                except OSError as e:
                    logger.warning("replace_file(%r) backup failed: %s", path, e)
                    return ActionResult(error=f"Failed to create backup {bak}: {e}")

            # 执行替换。
            if use_re:
                repl = new if regex else _literal_replacer(new)
                try:
                    if count is None:
                        new_content, replaced = pattern.subn(repl, content)
                    else:
                        new_content, replaced = pattern.subn(repl, content, count=count)
                except re.error as e:
                    # 替换模板（new 含非法 \g<...>）也会抛 re.error。
                    logger.warning("replace_file(%r) substitution failed: %s", path, e)
                    return ActionResult(error=f"Regex substitution failed for {old!r}: {e}")
            else:
                if count is None:
                    new_content = content.replace(old, new)
                    replaced = raw_total
                else:
                    new_content = content.replace(old, new, count)
                    replaced = min(count, raw_total)

            with open(tmp, "w", encoding=enc, newline=newline) as f:
                f.write(new_content)
            os.replace(tmp, path)
        except FileNotFoundError:
            return ActionResult(error=f"File not found: {path}")
        except UnicodeDecodeError as e:
            # 读取侧特有（write_file 是写入不会有）：文件不是合法 enc 编码。
            logger.warning("replace_file(%r) decode failed: %s", path, e)
            return ActionResult(error=f"Failed to decode {path} as {enc}: {e}")
        except LookupError as e:
            # 非法 encoding 名（非 OSError 子类）：单独兜底。
            if os.path.exists(tmp):
                try:
                    os.remove(tmp)
                except OSError:
                    pass
            logger.warning("replace_file(%r) unknown encoding %r: %s", path, enc, e)
            return ActionResult(error=f"Unknown encoding {enc!r}: {e}")
        except OSError as e:
            # 仅写阶段产生 tmp 残骸；读阶段失败时 tmp 不存在，exists 守卫保证安全。
            # 分级错误：权限/目录(path 指向目录)/磁盘满/只读 → 明确 error + warning，
            # 不冒泡到 Tools.execute 通用 catch（actions.py:260-262）。
            if os.path.exists(tmp):
                try:
                    os.remove(tmp)
                except OSError:
                    pass
            logger.warning("replace_file(%r) failed: %s", path, e)
            return ActionResult(error=f"Failed to replace text in {path}: {e}")

        # 回显：有限替换且未耗尽全部时注明 "of {raw_total}"；否则沿用原文案。
        # final_bytes 用替换后的 new_content（不是替换前的 content）。
        final_bytes = len(new_content.encode(enc))
        if count is not None and replaced < raw_total:
            match_clause = f"{replaced} of {raw_total} occurrence{'s' if raw_total != 1 else ''}"
        else:
            match_clause = f"{replaced} occurrence{'s' if replaced != 1 else ''}"
        memory = (
            f"Replaced {match_clause} of {old!r} with {new!r} "
            f"in {path} ({final_bytes} bytes)"
        )
        logger.info(memory)
        return ActionResult(extracted_content=memory, long_term_memory=memory)

    async def _action_evaluate(self, params: dict, browser: BrowserSession) -> ActionResult:
        code = params["code"]
        # 阶段二（二.B–二.F）：新增参数。registry 不校验 execute 路径（params 是 raw
        # dict），故在此读取 + 运行时守卫；见 memory action-params-no-runtime-validation。
        await_promise = params.get("await_promise", True)
        timeout_ms = params.get("timeout_ms")
        user_gesture = params.get("user_gesture", False)
        args = params.get("args")
        elements = params.get("elements")
        return_element_ids = params.get("return_element_ids", False)
        frame = params.get("frame")
        extract_images = params.get("extract_images", False)
        if timeout_ms is not None and not (1 <= timeout_ms <= 300000):
            return ActionResult(
                error=f"Evaluate failed: timeout_ms must be in [1, 300000], got {timeout_ms}",
            )
        if args is not None:
            try:
                json.dumps(args)  # 进 CDP 前先验可序列化
            except (TypeError, ValueError) as e:
                return ActionResult(error=f"Evaluate failed: args not JSON-serializable: {e}")
        if elements is not None and (
            not isinstance(elements, list) or not all(isinstance(i, int) for i in elements)
        ):
            return ActionResult(
                error="Evaluate failed: elements must be a list of ints (backend node ids)",
            )
        try:
            text = await browser.evaluate(
                code,
                args=args,
                elements=elements,
                await_promise=await_promise,
                timeout_ms=timeout_ms,
                user_gesture=user_gesture,
                return_element_ids=return_element_ids,
                frame=frame,
            )
        except Exception as e:
            # Hard error: JS exception / wasThrown / CDP failure — surface an
            # evaluate-specific error rather than the generic Tools.execute
            # fallback. Aligns with _action_find_elements / _action_search_page.
            logger.warning("evaluate(%r) failed: %s", code[:120], e)
            return ActionResult(error=f"Evaluate failed: {e}")
        # 二.D OUT: session 把返回节点解析回 "backendNodeId:<id>" → 作为可操作 index 回显
        if text.startswith("backendNodeId:"):
            bid = text.split(":", 1)[1]
            visible = (f"Returned element backend node id: {bid} "
                       f"(usable as index/element_id for click/input_text; "
                       f"if not in current selector_map, call get_state to refresh)")
            return ActionResult(extracted_content=visible,
                                long_term_memory=f"evaluate returned element index {bid}")
        # 二.F: 抽出 base64 图片，避免撑爆上下文
        metadata = None
        if extract_images:
            text, images = _extract_data_images(text)
            if images:
                metadata = {"images": images}
        # 二.A: 大结果分级落盘（镜像 _action_find_elements:1122-1144；OSError 不失败只 warning）
        tr = self._truncation
        saved_to = None
        if len(text) >= tr.eval_save_threshold:
            try:
                os.makedirs(tr.eval_output_dir, exist_ok=True)
                fpath = os.path.join(tr.eval_output_dir, f"evaluate_{int(time.time() * 1000)}.txt")
                with open(fpath, "w", encoding="utf-8", newline="") as f:
                    f.write(text)
                saved_to = fpath
            except OSError as e:
                logger.warning("evaluate: save to file failed: %s", e)
        limit = tr.eval_result_max_chars
        if saved_to:
            visible = (f"Evaluate result ({len(text)} chars) saved to {saved_to}. "
                       f"Preview: {text[:200]}...").strip()
            memory = f"JavaScript executed successfully, result saved: {saved_to}"
        else:
            visible = text[:limit]
            memory = _eval_long_term_memory(text)
        return ActionResult(extracted_content=visible, long_term_memory=memory, metadata=metadata)

    async def _action_search_page(self, params: dict, browser: BrowserSession) -> ActionResult:
        query = params["query"]
        try:
            data = await browser.search_page(
                query,
                regex=params.get("regex", False),
                case_sensitive=params.get("case_sensitive", False),
                context_chars=params.get("context_chars", 150),
                css_scope=params.get("css_scope"),
                max_results=params.get("max_results", 25),
                offset=params.get("offset", 0),
                search_attributes=params.get("search_attributes", False),
            )
        except Exception as e:
            # Hard error: CDP failure / invalid regex / css_scope not found —
            # surface a search_page-specific error rather than bubbling to the
            # generic Tools.execute catch (mirrors _action_find_text).
            logger.warning("search_page(%r) failed: %s", query, e)
            return ActionResult(error=f"Search page failed: {e}")
        total = data.get("total", 0)
        attr_total = data.get("attribute_total", 0)
        if total == 0 and not attr_total:
            # Soft miss: "text not on the page" is actionable info (the LLM can
            # scroll / switch tab / accept the text is absent), not a tool
            # failure — aligns with browser-use and _action_find_text (both
            # return extracted_content, not error, on a miss).
            msg = f"No matches for '{query}'"
            logger.info(msg)
            return ActionResult(extracted_content=msg, long_term_memory=msg)
        formatted = _format_search_results(data, query)
        # 大结果分级落盘（镜像 _action_extract；OSError 不失败只 warning）
        tr = self._truncation
        saved_to = None
        if len(formatted) >= tr.search_page_save_threshold:
            try:
                os.makedirs(tr.search_page_output_dir, exist_ok=True)
                fpath = os.path.join(tr.search_page_output_dir, f"search_page_{int(time.time() * 1000)}.txt")
                with open(fpath, "w", encoding="utf-8", newline="") as f:
                    f.write(formatted)
                saved_to = fpath
            except OSError as e:
                logger.warning("search_page: save to file failed: %s", e)
        if saved_to:
            visible = (f"Search results ({len(formatted)} chars) saved to {saved_to}. "
                       f"Preview: {formatted[:200]}...").strip()
        else:
            visible = formatted
        memory = f'Searched page for "{query}": {total} match{"es" if total != 1 else ""} found.'
        if attr_total:
            memory += f' (+{attr_total} attribute match{"es" if attr_total != 1 else ""})'
        if saved_to:
            memory += f" Results saved: {saved_to}"
        logger.info(memory)
        return ActionResult(extracted_content=visible, long_term_memory=memory)

    async def _action_done(self, params: dict, browser: BrowserSession) -> ActionResult:
        success = params.get("success", True)

        # 二.B 共享：解析 files_to_display → attachments（白名单 + 存在性）。
        # done 必须终止，故任何失败只 warn + 跳过，绝不 error / 绝不 is_done=False。
        attachments: list[str] = []
        allowed = self._allowed_read_paths
        for raw in params.get("files_to_display") or []:
            p = os.path.abspath(raw)
            if allowed and not any(p.startswith(pre) for pre in allowed):
                logger.warning("done: skip attachment outside allowed_read_paths: %s", p)
                continue
            if not os.path.isfile(p):
                logger.warning("done: skip missing attachment: %s", p)
                continue
            attachments.append(p)

        # 二.E 变体 B：结构化输出（output_model 给定时）。extracted_content 保持纯 JSON，
        # 不追加清单 / 不内联（避免破坏结构化输出）；原始数据另存 metadata。
        if self._output_model is not None:
            try:
                data = self._output_model.model_validate(params["data"])
            except (ValidationError, KeyError, TypeError) as e:
                # done 必须终止：结构化校验失败 → success=False 兜底，仍 is_done=True
                logger.warning("done: structured data invalid: %s", e)
                memory = "Task completed: False - invalid structured output"
                logger.info(memory)
                return ActionResult(
                    is_done=True,
                    success=False,
                    extracted_content=f"(invalid structured output: {e})",
                    long_term_memory=memory,
                    attachments=attachments or None,
                )
            payload = data.model_dump(mode="json")
            memory = f"Task completed (structured): {success}"
            logger.info(memory)
            return ActionResult(
                is_done=True,
                success=success,
                extracted_content=json.dumps(payload, ensure_ascii=False, indent=2),
                long_term_memory=memory,
                metadata={"structured_output": payload},
                attachments=attachments or None,
            )

        # 变体 A：自由文本（阶段一 + 二.A 后缀 + 二.B 清单 + 二.D 内联）
        text = (params.get("text") or "").strip()
        if not text:
            # done 必须终止（is_done=True 才退出循环，step.py:103），空 text 不能走
            # soft-miss（会变非终止循环）。兜底默认值保证终止 + 让退化情形在日志可见。
            text = "(no summary provided)"
            logger.warning("done called with empty text; substituting default summary")
        truncated = text[:100]
        memory = f"Task completed: {success} - {truncated}"
        if len(text) > 100:  # 二.A：对齐 browser-use 变体 A 的 - N more characters 后缀
            memory += f" - {len(text) - 100} more characters"
        visible = text
        if attachments:  # 二.B：附件清单（让 final_result() 可见；__str__ 仍被 500 截断）
            visible += "\n\nAttachments: " + ", ".join(os.path.basename(a) for a in attachments)
        if self._display_files_in_done_text and attachments:  # 二.D：内联文件内容
            cap = self._truncation.done_attachment_max_chars
            inline_parts = ["", "Attachments:"]
            for a in attachments:
                try:
                    with open(a, "r", encoding="utf-8", errors="replace") as f:
                        body = f.read(cap)
                except OSError as e:
                    logger.warning("done: skip inline read %s: %s", a, e)
                    continue
                inline_parts.append(f"--- {a} ---\n{body}")
            if len(inline_parts) > 1:
                visible += "\n" + "\n".join(inline_parts)
        logger.info(memory)
        return ActionResult(
            is_done=True,
            success=success,
            extracted_content=visible,
            long_term_memory=memory,
            attachments=attachments or None,
        )

    # ── Helpers ────────────────────────────────────────────────────────

    def _flatten_params(self, params: dict, action_name: str) -> dict:
        """Handle nested params like {"click": {"index": 5}} -> {"index": 5}.

        Won't unwrap a single nested dict whose key is a real field of this
        action's param_model (e.g. variant-B done ``data``) — that's a
        legitimate dict-valued param, not LLM wrapping.
        """
        if not params:
            return params
        if action_name in params and isinstance(params[action_name], dict):
            return params[action_name]
        # Check for any single nested dict value (common LLM pattern)
        dict_vals = {k: v for k, v in params.items() if isinstance(v, dict)}
        if len(dict_vals) == 1 and len(params) == 1:
            only_key = next(iter(dict_vals))
            param_model = self.registry.actions[action_name].param_model
            # 二.E：若该键是动作参数模型的真字段（如 done 变体 B 的 data），
            # 则是合法 dict 值参数，而非 LLM 包裹，不展开。
            if only_key not in param_model.model_fields:
                return list(dict_vals.values())[0]
        return params

    @staticmethod
    def _normalize(result: ActionResult | str | None) -> ActionResult:
        if isinstance(result, ActionResult):
            return result
        if isinstance(result, str):
            return ActionResult(extracted_content=result)
        return ActionResult()
